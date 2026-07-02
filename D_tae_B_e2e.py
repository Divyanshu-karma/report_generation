#!/usr/bin/env python3
from __future__ import annotations

import argparse
import asyncio
import json
import logging
import os
import re
import sys
import traceback
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import Workbook

try:
    from sqlalchemy import DateTime
    from sqlalchemy import Float
    from sqlalchemy import ForeignKey
    from sqlalchemy import Integer
    from sqlalchemy import String
    from sqlalchemy import Text
    from sqlalchemy import create_engine
    from sqlalchemy import func
    from sqlalchemy import select
    from sqlalchemy.orm import DeclarativeBase
    from sqlalchemy.orm import Mapped
    from sqlalchemy.orm import mapped_column
    from sqlalchemy.orm import relationship
    from sqlalchemy.orm import selectinload
    from sqlalchemy.orm import sessionmaker
except ImportError as exc:  # pragma: no cover - explicit runtime guard
    raise SystemExit(
        "SQLAlchemy is required for TAE_B_E2E_PDF_RUNNER_FOR_SQLA.py. "
        "Install it with: pip install SQLAlchemy"
    ) from exc


logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[logging.StreamHandler(sys.stdout)],
)
logger = logging.getLogger("TAE_B_E2E_PDF_RUNNER_FOR_SQLA")

_BOSS_PATH = Path(r"D:/docu3cProject/RG-AldrinP")
ENGINE_BETA = _BOSS_PATH.resolve() if _BOSS_PATH.exists() else Path(__file__).parent.resolve()
OUTPUT_DIR = ENGINE_BETA / "TAEB_Calib"
RUN_ARTIFACTS_ROOT = ENGINE_BETA / "temp" / "TAE_B_E2E_runs"
DATABASE_URL = os.getenv("DATABASE_URL")
if not DATABASE_URL:
    DATABASE_URL = f"sqlite:///{(OUTPUT_DIR / 'calibration.db').as_posix()}"
ENGINE_VERSION = "TM_BETA_TEST_LOCAL"


def _artifact_slug(value: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9._-]+", "_", (value or "").strip())
    slug = slug.strip("._-")
    return slug or "run"


def _build_run_artifacts_dir(pdf_paths: List[Path], benchmark_name: Optional[str]) -> Path:
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S_%f")
    first_pdf = pdf_paths[0].stem if pdf_paths else "run"
    bench_slug = _artifact_slug(benchmark_name or first_pdf)
    run_dir = RUN_ARTIFACTS_ROOT / f"{stamp}_{bench_slug}"
    run_dir.mkdir(parents=True, exist_ok=True)
    return run_dir


def _add_run_file_handler(log_path: Path) -> logging.Handler:
    handler = logging.FileHandler(log_path, encoding="utf-8")
    handler.setLevel(logging.INFO)
    handler.setFormatter(logging.Formatter("%(asctime)s [%(levelname)s] %(message)s"))
    logger.addHandler(handler)
    return handler



def _sanitize_tm_full_json_trace(match: Dict[str, Any]) -> str:
    trace = dict(match)
    trace.pop("image_base64", None)
    trace.pop("Image_Base64", None)
    trace.pop("base64", None)
    trace.pop("image", None)
    return json.dumps(trace, ensure_ascii=False, indent=2)


def _write_trace_artifact_to_xlsx_dir(output_xlsx: Path, index: int, match: Dict[str, Any], trace_json: str) -> None:
    artifact_dir = output_xlsx.parent / f"{output_xlsx.stem}_tm_full_json_traces"
    artifact_dir.mkdir(parents=True, exist_ok=True)
    trademark = str(match.get("Trademark", "") or match.get("TM_name", "") or "trace")
    slug = re.sub(r"[^A-Za-z0-9._-]+", "_", trademark).strip("_")
    if not slug:
        slug = "trace"
    filename = f"{index:05d}_{slug}.json"
    filepath = artifact_dir / filename
    filepath.write_text(trace_json, encoding="utf-8")


TRACE_CHUNK_SIZE = 32000


def _split_trace_payload(trace_json: str) -> tuple[str, str]:
    if not trace_json:
        return "", ""
    try:
        trace_payload = json.loads(trace_json)
    except Exception:
        return trace_json, ""

    score_ledger = trace_payload.pop("Score_Ledger", [])
    cleaned_trace_json = json.dumps(trace_payload, ensure_ascii=False, indent=2)
    score_ledger_json = json.dumps(score_ledger, ensure_ascii=False, indent=2)
    return cleaned_trace_json, score_ledger_json


class Base(DeclarativeBase):
    pass


class CalibrationRun(Base):
    __tablename__ = "calibration_runs_ROR"

    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), nullable=False, default=lambda: datetime.now(timezone.utc))
    pdf_path: Mapped[str] = mapped_column(String(1024), nullable=False)
    pdf_name: Mapped[str] = mapped_column(String(255), nullable=False, index=True)

    # NEW — required for the regression gate
    benchmark_name: Mapped[Optional[str]] = mapped_column(String(100), nullable=True, index=True)
    run_label: Mapped[Optional[str]] = mapped_column(String(50), nullable=True)
    git_commit: Mapped[Optional[str]] = mapped_column(String(40), nullable=True, index=True)

    engine_name: Mapped[str] = mapped_column(String(255), nullable=False)
    engine_version: Mapped[str] = mapped_column(String(255), nullable=False, default=ENGINE_VERSION)
    success: Mapped[bool] = mapped_column(nullable=False)
    error_message: Mapped[Optional[str]] = mapped_column(Text, nullable=True)
    traceback: Mapped[Optional[str]] = mapped_column(Text, nullable=True)
    match_count: Mapped[int] = mapped_column(Integer, nullable=False, default=0)

    # NEW — per-run score against ground truth
    correct_count: Mapped[Optional[int]] = mapped_column(Integer, nullable=True)
    ground_truth_count: Mapped[Optional[int]] = mapped_column(Integer, nullable=True)
    score: Mapped[Optional[float]] = mapped_column(Float, nullable=True)

    matches: Mapped[List["TrademarkMatch"]] = relationship(
        back_populates="run", cascade="all, delete-orphan", lazy="selectin",
    )


class TrademarkMatch(Base):
    __tablename__ = "trademark_matches_ROR"

    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    run_id: Mapped[int] = mapped_column(ForeignKey("calibration_runs_ROR.id", ondelete="CASCADE"), nullable=False, index=True)
    trademark: Mapped[str] = mapped_column(String(512), nullable=False, default="")
    owner: Mapped[str] = mapped_column(String(512), nullable=False, default="")
    serial_number: Mapped[Optional[str]] = mapped_column(String(128), nullable=True)
    registration_number: Mapped[Optional[str]] = mapped_column(String(128), nullable=True)
    match_key: Mapped[str] = mapped_column(String(255), nullable=False, default="", index=True)  # NEW
    score: Mapped[float] = mapped_column(Float, nullable=False, default=0.0)
    risk_level: Mapped[str] = mapped_column(String(64), nullable=False, default="CLEARED")
    status: Mapped[str] = mapped_column(String(64), nullable=False, default="")
    reasoning: Mapped[str] = mapped_column(Text, nullable=False, default="")
    tm_full_json_trace: Mapped[str] = mapped_column(Text, nullable=False, default="")

    # NEW — verdict against ground truth
    expected_level: Mapped[Optional[str]] = mapped_column(String(20), nullable=True)
    is_correct: Mapped[Optional[bool]] = mapped_column(nullable=True)
    verdict: Mapped[Optional[str]] = mapped_column(String(10), nullable=True)  # "✓" / "↑" / "↓"

    # NEW — promoted trace fields, extracted at ingestion (Gap 5)
    visual_sim: Mapped[Optional[float]] = mapped_column(Float, nullable=True)
    goods_overlap: Mapped[Optional[float]] = mapped_column(Float, nullable=True)
    llm_factor: Mapped[Optional[str]] = mapped_column(String(50), nullable=True)
    llm_gs: Mapped[Optional[float]] = mapped_column(Float, nullable=True)
    pre_llm_score: Mapped[Optional[float]] = mapped_column(Float, nullable=True)

    run: Mapped[CalibrationRun] = relationship(back_populates="matches")


class GroundTruthAttorney(Base):
    __tablename__ = "ground_truth_attorney_ROR"

    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)

    # CHANGED — keyed by benchmark, not pdf_name
    benchmark_name: Mapped[str] = mapped_column(String(100), nullable=False, index=True)
    pdf_name: Mapped[Optional[str]] = mapped_column(String(255), nullable=True)  # retained for traceability only

    tm_name: Mapped[str] = mapped_column(String(512), nullable=False)
    serial_number: Mapped[Optional[str]] = mapped_column(String(128), nullable=True)  # NEW — for match_key join
    registration_number: Mapped[Optional[str]] = mapped_column(String(128), nullable=True)  # NEW
    owner_name: Mapped[Optional[str]] = mapped_column(String(512), nullable=True)

    # CHANGED — normalized to RiskLevel enum values at write time
    expected_level: Mapped[str] = mapped_column(String(20), nullable=False)
    attorney_verdict_reasoning: Mapped[Optional[str]] = mapped_column(Text, nullable=True)
    validated_by: Mapped[Optional[str]] = mapped_column(String(100), nullable=True)  # NEW
    validated_at: Mapped[Optional[datetime]] = mapped_column(DateTime(timezone=True), nullable=True)  # NEW
    is_boundary_case: Mapped[bool] = mapped_column(nullable=False, default=False)  # NEW

    @property
    def risk_of_registration(self) -> Optional[str]:
        return self.expected_level

    @risk_of_registration.setter
    def risk_of_registration(self, value: Optional[str]) -> None:
        self.expected_level = value or "CLEARED"




ENGINE = create_engine(DATABASE_URL, future=True)
SessionLocal = sessionmaker(bind=ENGINE, autoflush=False, autocommit=False, expire_on_commit=False)


def bootstrap_database() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    Base.metadata.create_all(ENGINE)


LEVEL_ORDER = ["HIGH", "MEDIUM_HIGH", "MEDIUM", "MEDIUM_LOW", "LOW", "CLEARED"]
ACCEPTABLE_TRANSITIONS = {
    ("LOW", "CLEARED"),
    ("LOW", "MEDIUM_LOW"),
    ("LOW", "MEDIUM"),
    ("CLEARED", "LOW"),
    ("CLEARED", "MEDIUM_LOW"),
    ("CLEARED", "MEDIUM"),
    ("MEDIUM_LOW", "LOW"),
    ("MEDIUM_LOW", "CLEARED"),
    ("MEDIUM_LOW", "MEDIUM"),
    ("MEDIUM", "LOW"),
    ("MEDIUM", "CLEARED"),
    ("MEDIUM", "MEDIUM_LOW"),
}
LEVEL_ALIASES = {
    "HIGH RISK": "HIGH",
    "MEDIUM-HIGH": "MEDIUM_HIGH",
    "MEDIUM HIGH": "MEDIUM_HIGH",
    "MED HIGH": "MEDIUM_HIGH",
    "MED-HIGH": "MEDIUM_HIGH",
    "MEDIUM RISK": "MEDIUM",
    "MEDIUM-LOW": "MEDIUM_LOW",
    "MEDIUM LOW": "MEDIUM_LOW",
    "MED LOW": "MEDIUM_LOW",
    "MED-LOW": "MEDIUM_LOW",
    "MED": "MEDIUM",
    "LOW RISK": "LOW",
    "CLEAR": "CLEARED",
    "NO RISK": "CLEARED",
    "CLEARED": "CLEARED",
}

def normalize_level(raw: str) -> Optional[str]:
    s = (raw or "").strip().upper()
    if s in LEVEL_ORDER:
        return s
    return LEVEL_ALIASES.get(s)


def levels_are_acceptable(expected: Optional[str], actual: Optional[str]) -> bool:
    if not expected or not actual:
        return False
    if expected == actual:
        return True
    return (expected, actual) in ACCEPTABLE_TRANSITIONS


def acceptance_label(expected: Optional[str], actual: Optional[str]) -> str:
    expected_norm = normalize_level(expected or "")
    actual_norm = normalize_level(actual or "")
    if not expected_norm or not actual_norm:
        return "No"
    if expected_norm == actual_norm:
        return "Yes"
    if levels_are_acceptable(expected_norm, actual_norm):
        return "acceptable deviation"
    return "No"


def save_ground_truth(benchmark_name: str, gt_entries: List[Dict[str, Any]], pdf_name: Optional[str] = None) -> None:
    from sqlalchemy import delete
    if pdf_name is None:
        pdf_name = benchmark_name
    with SessionLocal() as session:
        session.execute(
            delete(GroundTruthAttorney).where(
                (GroundTruthAttorney.benchmark_name == benchmark_name) |
                (GroundTruthAttorney.pdf_name == pdf_name)
            )
        )
        for entry in gt_entries:
            raw_tm = str(entry.get("tm_name") or "").strip()
            # Split by '/'
            parts = [p.strip() for p in raw_tm.split("/") if p.strip()]
            for part in parts:
                # Remove pattern like (x2) or (×2) or (x 2)
                cleaned = re.sub(r'\s*\(\s*[xX×*]?\s*\d+\s*\)', '', part, flags=re.IGNORECASE)
                cleaned = cleaned.strip()
                if not cleaned:
                    continue

                raw_risk = entry.get("expected_level") or entry.get("risk_of_registration") or "CLEARED"
                expected = normalize_level(raw_risk)
                if expected is None:
                    raise ValueError(
                        f"Ground truth risk level '{raw_risk}' fails to normalize. "
                        f"Expected one of {LEVEL_ORDER} or known aliases."
                    )

                validated_at_val = entry.get("validated_at")
                if isinstance(validated_at_val, str):
                    validated_at_val = datetime.fromisoformat(validated_at_val.replace("Z", "+00:00"))

                gt_record = GroundTruthAttorney(
                    benchmark_name=benchmark_name,
                    pdf_name=pdf_name,
                    tm_name=cleaned,
                    serial_number=entry.get("serial_number") or entry.get("Serial_Number"),
                    registration_number=entry.get("registration_number") or entry.get("Registration_Number"),
                    owner_name=entry.get("owner_name") or entry.get("Owner"),
                    expected_level=expected,
                    attorney_verdict_reasoning=entry.get("attorney_verdict_reasoning") or entry.get("Reasoning"),
                    validated_by=entry.get("validated_by"),
                    validated_at=validated_at_val,
                    is_boundary_case=bool(entry.get("is_boundary_case", False)),
                )
                session.add(gt_record)
        session.commit()


def load_ground_truth(key: str) -> List[GroundTruthAttorney]:
    with SessionLocal() as session:
        stmt = select(GroundTruthAttorney).where(
            (GroundTruthAttorney.benchmark_name == key) |
            (GroundTruthAttorney.pdf_name == key)
        )
        return list(session.execute(stmt).scalars().all())



RUNNER_SCRIPT_CONTENT = """
import sys
import os
import json
import asyncio
import re
import traceback
from pathlib import Path
from typing import Any, Dict

engine_dir = sys.argv[1]
pdf_path = sys.argv[2]
output_json = sys.argv[3]

sys.path.insert(0, str(Path(engine_dir).resolve()))

os.environ["CUDA_VISIBLE_DEVICES"] = ""
os.environ["JSON_FILE"] = "True"

phoenix_disabled = os.getenv("PHOENIX_DISABLED", "false").lower() in {"1", "true", "yes"}
if phoenix_disabled:
    os.environ["PHOENIX_DISABLED"] = "true"
else:
    os.environ.pop("PHOENIX_DISABLED", None)

from src.main import run_pipeline
from src.attorney_view import get_attorney_view_matches

try:
    from phoenix.otel import register as phoenix_register
except Exception:
    phoenix_register = None

try:
    from opentelemetry import trace as otel_trace
except Exception:
    otel_trace = None


def sanitize_tm_full_json_trace(match):
    trace = dict(match)
    trace.pop("image_base64", None)
    trace.pop("Image_Base64", None)
    trace.pop("base64", None)
    trace.pop("image", None)
    return json.dumps(trace, ensure_ascii=False, indent=2)


def _trace_artifact_directory(output_xlsx: Path) -> Path:
    return output_xlsx.parent / f"{output_xlsx.stem}_tm_full_json_traces"


def _trace_artifact_name(index: int, match: Dict[str, Any]) -> str:
    trademark = str(match.get("Trademark", "") or match.get("TM_name", "") or "trace")
    slug = re.sub(r"[^A-Za-z0-9._-]+", "_", trademark).strip("_")
    if not slug:
        slug = "trace"
    return f"{index:05d}_{slug}.json"


def _write_trace_artifact(output_xlsx: Path, index: int, match: Dict[str, Any], trace_json: str) -> str:
    artifact_dir = _trace_artifact_directory(output_xlsx)
    artifact_dir.mkdir(parents=True, exist_ok=True)
    artifact_path = artifact_dir / _trace_artifact_name(index, match)
    artifact_path.write_text(trace_json, encoding="utf-8")
    return str(artifact_path.relative_to(output_xlsx.parent))


async def main():
    tracer_provider = None
    tracer = None
    span = None
    if phoenix_register is not None and not phoenix_disabled:
        try:
            tracer_provider = phoenix_register(project_name="TM_BETA_TEST_LOCAL")
            if otel_trace is not None:
                tracer = otel_trace.get_tracer(__name__)
                span = tracer.start_span("TAE_B_E2E_PDF_RUNNER")
                span.__enter__()
        except Exception:
            tracer_provider = None
            tracer = None
            span = None

    try:
        if span is not None:
            span.set_attribute("pdf_path", pdf_path)
        ext = os.path.splitext(pdf_path)[1].lower()
        report_type = "fovea_docx" if ext == ".docx" else "uspto"
        result = await run_pipeline(pdf_path, report_type, skip_reasoning=False)
        query_app = result.get("query_app")

        tm_name = getattr(query_app, "mark_text", None) or ""
        tm_owner = getattr(query_app, "applicant_name", None) or ""

        matches = get_attorney_view_matches(result)
        if not matches:
            matches = []
            for citation in result.get("citations", []) or []:
                matches.append({
                    "Trademark": citation.get("Trademark", ""),
                    "Owner": citation.get("Owner", ""),
                    "Serial_Number": citation.get("Serial_Number"),
                    "Registration_Number": citation.get("Registration_Number"),
                    "Score": citation.get("Confusion_Risk_Score", citation.get("Score", 0.0)),
                    "Risk_Level": citation.get("Risk_Level", "CLEARED"),
                    "Reasoning": citation.get(
                        "Reasoning",
                        "Included from parsed citations because attorney-view matches were not produced.",
                    ),
                    "Class": citation.get("Class", ""),
                    "Status": citation.get("Status", ""),
                })

        serialized_matches = []
        for m in matches:
            serialized_matches.append({
                "Trademark": m.get("Trademark", ""),
                "Owner": m.get("Owner", ""),
                "Serial_Number": m.get("Serial_Number"),
                "Registration_Number": m.get("Registration_Number"),
                "Score": m.get("Score", 0.0),
                "Risk_Level": m.get("Risk_Level", "CLEARED"),
                "Reasoning": m.get("Reasoning", ""),
                "Class": m.get("Class", ""),
                "Status": m.get("Status", ""),
                "TM_FULL_JSON_TRACE": sanitize_tm_full_json_trace(m),
            })

        out_data = {
            "success": True,
            "tm_name": tm_name,
            "tm_owner": tm_owner,
            "matches": serialized_matches
        }
    except Exception as e:
        out_data = {
            "success": False,
            "error": str(e),
            "traceback": traceback.format_exc()
        }
    finally:
        if span is not None:
            try:
                span.__exit__(None, None, None)
            except Exception:
                pass
        if tracer_provider is not None:
            try:
                tracer_provider.shutdown()
            except Exception:
                pass

    with open(output_json, "w", encoding="utf-8") as f:
        json.dump(out_data, f, indent=2)


if __name__ == "__main__":
    asyncio.run(main())
"""


async def run_single_engine(
    runner_path: Path, engine_dir: Path, pdf_path: Path, temp_dir: Path
) -> Dict[str, Any]:
    out_json_path = temp_dir / f"result_{engine_dir.name}_{pdf_path.stem}.json"
    cmd = [
        sys.executable,
        str(runner_path),
        str(engine_dir),
        str(pdf_path),
        str(out_json_path),
    ]

    try:
        proc = await asyncio.create_subprocess_exec(
            *cmd,
            cwd=str(engine_dir),
            stdout=asyncio.subprocess.PIPE,
            stderr=asyncio.subprocess.STDOUT,
            env={**os.environ, "PYTHONUNBUFFERED": "1"},
        )

        if proc.stdout is not None:
            while True:
                raw_line = await proc.stdout.readline()
                if not raw_line:
                    break
                line = raw_line.decode("utf-8", errors="replace").rstrip()
                if line:
                    logger.info("[USPTO_ENGINE:%s] %s", pdf_path.name, line)

        await proc.wait()

        if proc.returncode != 0:
            logger.error("Engine %s failed for %s (code %s)", engine_dir.name, pdf_path.name, proc.returncode)
            return {"success": False, "error": f"Process exited with non-zero code {proc.returncode}."}

        if not out_json_path.exists():
            return {"success": False, "error": "Engine completed but output JSON was not generated."}

        with open(out_json_path, "r", encoding="utf-8") as f:
            data = json.load(f)

        if not data.get("success"):
            logger.error(
                "Engine %s reported failure for %s: %s",
                engine_dir.name,
                pdf_path.name,
                data.get("error") or "AssertionError / Unknown Error",
            )
            if data.get("traceback"):
                logger.error("Engine failure traceback:\n%s", data.get("traceback"))

        return data
    except Exception as e:
        logger.error("Error running engine %s for %s: %s", engine_dir.name, pdf_path.name, e)
        return {"success": False, "error": str(e), "traceback": traceback.format_exc()}


def get_match_key(match: Dict[str, Any]) -> str:
    name = str(match.get("Trademark") or "").strip().lower()
    serial = str(match.get("Serial_Number") or "").strip()
    reg = str(match.get("Registration_Number") or "").strip()
    cls = str(match.get("Class") or "").strip()

    name_clean = "".join(c for c in name if c.isalnum())
    serial_clean = "".join(c for c in serial if c.isdigit())
    reg_clean = "".join(c for c in reg if c.isdigit())

    if serial_clean and serial_clean != "00000000":
        return f"sn_{serial_clean}"
    if reg_clean:
        return f"rn_{reg_clean}"
    return f"comp_{name_clean}_{cls}"


def _extract_trace_fields(trace_json: str) -> Dict[str, Any]:
    try:
        t = json.loads(trace_json)
    except Exception:
        return {}
    sim = t.get("Similarity_Analysis") or {}
    ledger = t.get("Score_Ledger") or []
    llm5d = t.get("LLM_5D_Results") or {}
    llm_factor = next((
        e.get("legal_factor") for e in ledger
        if isinstance(e, dict) and e.get("module_name") == "LLMValidation"
        and e.get("legal_factor") in (
            "llm_batch_confirmation", "llm_partial_downgrade",
            "llm_goods_channel_downgrade", "llm_batch_neutral",
            "llm_batch_neutral_ceiling",
        )
    ), None)
    return {
        "visual_sim": float(sim.get("Visual") or 0) if sim.get("Visual") is not None else None,
        "goods_overlap": float(sim.get("Goods_Overlap") or 0) if sim.get("Goods_Overlap") is not None else None,
        "llm_factor": llm_factor,
        "llm_gs": float((llm5d.get("V2_Goods_Score") or 0)) / 100.0 if (llm5d and llm5d.get("V2_Goods_Score") is not None) else None,
        "pre_llm_score": float(t.get("pre_llm_risk_score")) if t.get("pre_llm_risk_score") is not None else None,
    }


def persist_calibration_run(
    *,
    pdf_path: Path,
    engine_name: str,
    result: Dict[str, Any],
    matches: List[Dict[str, Any]],
    benchmark_name: Optional[str] = None,
    run_label: Optional[str] = None,
    git_commit: Optional[str] = None,
) -> int:
    success = bool(result.get("success"))
    error_message = None if success else str(result.get("error") or "Unknown failure")
    tb = None if success else result.get("traceback") or traceback.format_exc()

    run = CalibrationRun(
        pdf_path=str(pdf_path),
        pdf_name=pdf_path.name,
        benchmark_name=benchmark_name,
        run_label=run_label,
        git_commit=git_commit,
        engine_name=engine_name,
        engine_version=ENGINE_VERSION,
        success=success,
        error_message=error_message,
        traceback=tb,
        match_count=len(matches),
    )

    if success:
        for match in matches:
            trace_json = str(match.get("TM_FULL_JSON_TRACE", ""))
            trace_fields = _extract_trace_fields(trace_json)
            run.matches.append(
                TrademarkMatch(
                    trademark=str(match.get("Trademark", "")),
                    owner=str(match.get("Owner", "")),
                    serial_number=match.get("Serial_Number"),
                    registration_number=match.get("Registration_Number"),
                    match_key=get_match_key(match),
                    score=float(match.get("Score", 0.0) or 0.0),
                    risk_level=str(match.get("Risk_Level", "CLEARED")),
                    status=str(match.get("Status", "")),
                    reasoning=str(match.get("Reasoning", "")),
                    tm_full_json_trace=trace_json,
                    visual_sim=trace_fields.get("visual_sim"),
                    goods_overlap=trace_fields.get("goods_overlap"),
                    llm_factor=trace_fields.get("llm_factor"),
                    llm_gs=trace_fields.get("llm_gs"),
                    pre_llm_score=trace_fields.get("pre_llm_score"),
                )
            )

    with SessionLocal() as session:
        session.add(run)
        session.commit()
        session.refresh(run)
        return run.id


def get_previous_runs(key: str) -> List[CalibrationRun]:
    with SessionLocal() as session:
        stmt = (
            select(CalibrationRun)
            .where(
                (CalibrationRun.pdf_name == key) |
                (CalibrationRun.benchmark_name == key)
            )
            .order_by(CalibrationRun.created_at.desc(), CalibrationRun.id.desc())
            .options(selectinload(CalibrationRun.matches))
        )
        return list(session.execute(stmt).scalars().all())


def compare_latest_two_runs(key: str) -> Dict[str, Any]:
    runs = get_previous_runs(key)
    latest = runs[0] if len(runs) > 0 else None
    previous = runs[1] if len(runs) > 1 else None

    def _shape(run: Optional[CalibrationRun]) -> Optional[Dict[str, Any]]:
        if run is None:
            return None
        return {
            "id": run.id,
            "created_at": run.created_at.isoformat() if run.created_at else None,
            "pdf_name": run.pdf_name,
            "success": run.success,
            "engine_version": run.engine_version,
            "match_count": run.match_count,
            "correct_count": run.correct_count,
            "ground_truth_count": run.ground_truth_count,
            "score": run.score,
        }

    return {"latest": _shape(latest), "previous": _shape(previous)}


def compute_verdicts(run_id: int, benchmark_name: str) -> None:
    """Join run's matches against ground truth, write expected_level/is_correct/verdict."""
    with SessionLocal() as session:
        # Get run info first to obtain the pdf_name
        run = session.get(CalibrationRun, run_id)
        if not run:
            logger.error("CalibrationRun with id %s not found for verdict computation", run_id)
            return
        pdf_name = run.pdf_name

        # Query ground truth by benchmark_name, or fallback to pdf_name (retaining compatibility)
        gt_rows = session.execute(
            select(GroundTruthAttorney).where(GroundTruthAttorney.benchmark_name == benchmark_name)
        ).scalars().all()

        if not gt_rows and pdf_name:
            gt_rows = session.execute(
                select(GroundTruthAttorney).where(
                    (GroundTruthAttorney.pdf_name == pdf_name) |
                    (GroundTruthAttorney.benchmark_name == pdf_name)
                )
            ).scalars().all()

        # Index GT by serial/reg number first, fall back to normalized name
        gt_by_key: Dict[str, GroundTruthAttorney] = {}
        gt_by_name: Dict[str, GroundTruthAttorney] = {}
        for gt in gt_rows:
            if gt.serial_number:
                gt_serial_clean = "".join(c for c in gt.serial_number if c.isdigit())
                if gt_serial_clean:
                    gt_by_key[f"sn_{gt_serial_clean}"] = gt
            if gt.registration_number:
                gt_reg_clean = "".join(c for c in gt.registration_number if c.isdigit())
                if gt_reg_clean:
                    gt_by_key[f"rn_{gt_reg_clean}"] = gt
            norm = "".join(c for c in gt.tm_name.lower() if c.isalnum())
            gt_by_name.setdefault(norm, gt)

        matches = session.execute(
            select(TrademarkMatch).where(TrademarkMatch.run_id == run_id)
        ).scalars().all()

        correct = 0
        total_gt = 0

        for m in matches:
            gt = gt_by_key.get(m.match_key)
            if gt is None:
                norm = "".join(c for c in m.trademark.lower() if c.isalnum())
                gt = gt_by_name.get(norm)
            if gt is None:
                continue

            total_gt += 1
            m.expected_level = gt.expected_level
            actual = (m.risk_level or "").strip().upper()

            if levels_are_acceptable(gt.expected_level, actual):
                m.is_correct = True
                m.verdict = "✓"
                correct += 1
            else:
                m.is_correct = False
                ai = LEVEL_ORDER.index(actual) if actual in LEVEL_ORDER else 9
                ei = LEVEL_ORDER.index(gt.expected_level) if gt.expected_level in LEVEL_ORDER else 9
                m.verdict = "↑" if ai < ei else "↓"

        run.ground_truth_count = total_gt
        run.correct_count = correct
        run.score = correct / total_gt if total_gt else None
        session.commit()


def scorecard_diff(benchmark_name: str, baseline_run_id: int, candidate_run_id: int) -> Dict[str, Any]:
    with SessionLocal() as session:
        def _results(run_id):
            rows = session.execute(
                select(TrademarkMatch).where(
                    TrademarkMatch.run_id == run_id,
                    TrademarkMatch.expected_level.isnot(None),
                )
            ).scalars().all()
            return {r.match_key: r for r in rows}

        a, b = _results(baseline_run_id), _results(candidate_run_id)
        fixed, regressed, stable_correct, stable_wrong = [], [], [], []

        for key, rb in b.items():
            ra = a.get(key)
            if ra is None:
                continue
            if not ra.is_correct and rb.is_correct:
                fixed.append(rb.trademark)
            elif ra.is_correct and not rb.is_correct:
                regressed.append(rb.trademark)
            elif rb.is_correct:
                stable_correct.append(rb.trademark)
            else:
                stable_wrong.append(rb.trademark)

        return {
            "benchmark": benchmark_name,
            "fixed": fixed,
            "regressed": regressed,
            "stable_correct": stable_correct,
            "stable_wrong": stable_wrong,
            "net_change": len(fixed) - len(regressed),
        }


def regression_gate(benchmark_name: str, baseline_run_id: int, candidate_run_id: int, allow_regressions: int = 0):
    diff = scorecard_diff(benchmark_name, baseline_run_id, candidate_run_id)
    passes = len(diff["regressed"]) <= allow_regressions
    return passes, diff




def get_run_statistics() -> Dict[str, Any]:
    with SessionLocal() as session:
        total_runs = session.scalar(select(func.count()).select_from(CalibrationRun)) or 0
        success_runs = session.scalar(select(func.count()).select_from(CalibrationRun).where(CalibrationRun.success.is_(True))) or 0
        failure_runs = session.scalar(select(func.count()).select_from(CalibrationRun).where(CalibrationRun.success.is_(False))) or 0
        total_matches = session.scalar(select(func.count()).select_from(TrademarkMatch)) or 0

        return {
            "total_runs": int(total_runs),
            "success_runs": int(success_runs),
            "failure_runs": int(failure_runs),
            "total_matches": int(total_matches),
        }


def _load_matches_for_export(run_id: int) -> List[Dict[str, Any]]:
    with SessionLocal() as session:
        stmt = select(TrademarkMatch).where(TrademarkMatch.run_id == run_id).order_by(TrademarkMatch.id.asc())
        rows = session.execute(stmt).scalars().all()
        return [
            {
                "Trademark": row.trademark,
                "Owner": row.owner,
                "Serial_Number": row.serial_number,
                "Registration_Number": row.registration_number,
                "Score": row.score,
                "Risk_Level": row.risk_level,
                "Reasoning": row.reasoning,
                "Status": row.status,
                "TM_FULL_JSON_TRACE": row.tm_full_json_trace,
                "Match_Key": row.match_key,
                "Expected_Level": row.expected_level,
            }
            for row in rows
        ]


async def process_pdf(
    runner_path: Path,
    pdf_path: Path,
    temp_dir: Path,
    benchmark_name: Optional[str] = None,
    run_label: Optional[str] = None,
    git_commit: Optional[str] = None,
) -> Dict[str, Any]:
    logger.info("[START] Processing PDF: %s", pdf_path.name)
    logger.info("--- RUNNING TM_BETA_TEST_LOCAL on %s ---", pdf_path.name)
    beta_res = await run_single_engine(runner_path, ENGINE_BETA, pdf_path, temp_dir)
    logger.info("[DONE] Completed running TM_BETA_TEST_LOCAL for PDF: %s", pdf_path.name)

    try:
        match_rows = []
        if beta_res.get("success"):
            match_rows = beta_res.get("matches", []) or []

        # Derive benchmark_name if not provided
        derived_benchmark = benchmark_name
        if not derived_benchmark:
            derived_benchmark = beta_res.get("tm_name")
        if not derived_benchmark:
            derived_benchmark = pdf_path.stem

        run_id = persist_calibration_run(
            pdf_path=pdf_path,
            engine_name="TM_BETA_TEST_LOCAL",
            result=beta_res,
            matches=match_rows,
            benchmark_name=derived_benchmark,
            run_label=run_label,
            git_commit=git_commit,
        )

        if beta_res.get("success") and run_id is not None:
            compute_verdicts(run_id, derived_benchmark)

        return {
            "file_path": str(pdf_path),
            "run_id": run_id,
            "beta": beta_res,
            "derived_benchmark": derived_benchmark,
        }
    except Exception as exc:
        logger.error("Failed to persist calibration run for %s: %s", pdf_path.name, exc)
        return {
            "file_path": str(pdf_path),
            "run_id": None,
            "beta": {
                "success": False,
                "error": str(exc),
                "traceback": traceback.format_exc(),
            },
            "derived_benchmark": benchmark_name or pdf_path.stem,
        }



async def main_async(
    pdf_paths: List[Path],
    output_xlsx: Path,
    gt_json: Optional[str] = None,
    gt_file: Optional[Path] = None,
    benchmark_name: Optional[str] = None,
    run_label: Optional[str] = None,
    git_commit: Optional[str] = None,
    allow_regressions: int = 0,
) -> int:
    bootstrap_database()
    output_xlsx.parent.mkdir(parents=True, exist_ok=True)

    gt_data = None
    if gt_json:
        try:
            gt_data = json.loads(gt_json)
        except Exception as exc:
            logger.error("Failed to parse --gt-json: %s", exc)
            return 1
    elif gt_file:
        try:
            resolved_gt_file = gt_file.resolve()
            if resolved_gt_file.exists():
                gt_data = json.loads(resolved_gt_file.read_text(encoding="utf-8"))
            else:
                logger.error("Ground Truth file not found: %s", gt_file)
                return 1
        except Exception as exc:
            logger.error("Failed to read --gt-file: %s", exc)
            return 1

    if gt_data is not None:
        if not isinstance(gt_data, list):
            logger.error("Ground Truth data must be a JSON array/list.")
            return 1
        for idx, entry in enumerate(gt_data):
            if not isinstance(entry, dict) or "tm_name" not in entry:
                logger.error("Ground Truth entry at index %s is invalid (must be a dict containing 'tm_name').", idx)
                return 1

    valid_pdfs = []
    for pdf_path in pdf_paths:
        resolved = pdf_path.resolve()
        if resolved.exists() and resolved.is_file():
            valid_pdfs.append(resolved)
            if gt_data is not None:
                bench = benchmark_name or resolved.stem
                save_ground_truth(bench, gt_data, pdf_name=resolved.name)
        else:
            logger.error("File not found or is invalid: %s (resolved: %s)", pdf_path, resolved)

    if not valid_pdfs:
        logger.error("No valid PDF/DOCX files found to benchmark.")
        return 1

    logger.info("Found %s valid target PDFs to benchmark.", len(valid_pdfs))

    temp_dir = _build_run_artifacts_dir(valid_pdfs, benchmark_name)
    log_path = temp_dir / f"{_artifact_slug(benchmark_name or valid_pdfs[0].stem)}_debug.log"
    file_handler = _add_run_file_handler(log_path)
    try:
        runner_path = temp_dir / "temp_runner.py"
        runner_path.write_text(RUNNER_SCRIPT_CONTENT, encoding="utf-8")

        results = []
        for pdf_path in valid_pdfs:
            results.append(
                await process_pdf(
                    runner_path,
                    pdf_path,
                    temp_dir,
                    benchmark_name=benchmark_name,
                    run_label=run_label,
                    git_commit=git_commit,
                )
            )

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Comparison"
        worksheet.append([
            "Acceptence",
            "TM_name",
            "TM_Owner",
            "Status",
            "TM_risk_score",
            "TM_Risk_level",
            "TM_reasoning",
            "TM_FULL_JSON_TRACE",
            "Score_Ledger_each_mark",
            "GT_TM_Name",
            "GT_Owner_Name",
            "GT_Risk_of_Registration",
            "GT_Attorney_Verdict_Reasoning",
        ])

        for res in results:
            file_path = res["file_path"]
            beta = res["beta"]
            run_id = res.get("run_id")
            derived_bench_name = res.get("derived_benchmark")

            if beta.get("success") and run_id is not None:
                gt_list = load_ground_truth(derived_bench_name)
                if not gt_list:
                    gt_list = load_ground_truth(Path(file_path).name)
                beta_matches = _load_matches_for_export(run_id)

                if gt_list:
                    def normalize_name(name: str) -> str:
                        if not name:
                            return ""
                        return "".join(c for c in name.strip().lower() if c.isalnum())

                    gt_by_key = {}
                    gt_by_name = {}
                    for gt in gt_list:
                        if gt.serial_number:
                            gt_serial_clean = "".join(c for c in gt.serial_number if c.isdigit())
                            if gt_serial_clean:
                                gt_by_key[f"sn_{gt_serial_clean}"] = gt
                        if gt.registration_number:
                            gt_reg_clean = "".join(c for c in gt.registration_number if c.isdigit())
                            if gt_reg_clean:
                                gt_by_key[f"rn_{gt_reg_clean}"] = gt
                        norm = normalize_name(gt.tm_name)
                        if norm:
                            gt_by_name.setdefault(norm, gt)

                    # Validation Layer: Track count of matched database items vs worksheet items
                    expected_matched_count = sum(1 for m in beta_matches if m.get("Expected_Level") is not None)
                    actual_matched_count = 0

                    for match_index, match in enumerate(beta_matches, start=1):
                        trace_json = str(match.get("TM_FULL_JSON_TRACE", ""))
                        _write_trace_artifact_to_xlsx_dir(output_xlsx, match_index, match, trace_json)
                        cleaned_trace_json, score_ledger_json = _split_trace_payload(trace_json)
                        match_tm = match.get("Trademark", "")
                        match_key = match.get("Match_Key")

                        gt = gt_by_key.get(match_key)
                        if gt is None:
                            norm_match = normalize_name(match_tm)
                            gt = gt_by_name.get(norm_match)

                        if gt is not None:
                            actual_matched_count += 1
                            acceptance = acceptance_label(
                                gt.expected_level,
                                match.get("Risk_Level", "CLEARED"),
                            )
                            worksheet.append([
                                acceptance,
                                match.get("Trademark", ""),
                                match.get("Owner", ""),
                                match.get("Status", ""),
                                match.get("Score", 0.0),
                                match.get("Risk_Level", "CLEARED"),
                                match.get("Reasoning", ""),
                                cleaned_trace_json,
                                score_ledger_json,
                                gt.tm_name,
                                gt.owner_name or "",
                                gt.expected_level,
                                gt.attorney_verdict_reasoning or "",
                            ])

                    if expected_matched_count != actual_matched_count:
                        raise ValueError(
                            f"Validation Layer Error: Run {run_id} has {expected_matched_count} matches in DB with ground truth verdicts, "
                            f"but only {actual_matched_count} matched and exported to XLSX comparison sheet. "
                            "Please check for inconsistencies between the Ground Truth names and Trademark names."
                        )
                else:
                    for match_index, match in enumerate(beta_matches, start=1):
                        trace_json = str(match.get("TM_FULL_JSON_TRACE", ""))
                        _write_trace_artifact_to_xlsx_dir(output_xlsx, match_index, match, trace_json)
                        cleaned_trace_json, score_ledger_json = _split_trace_payload(trace_json)
                        worksheet.append([
                            "No",
                            match.get("Trademark", ""),
                            match.get("Owner", ""),
                            match.get("Status", ""),
                            match.get("Score", 0.0),
                            match.get("Risk_Level", "CLEARED"),
                            match.get("Reasoning", ""),
                            cleaned_trace_json,
                            score_ledger_json,
                            "",
                            "",
                            "",
                            "",
                        ])
            else:
                failure_reason = beta.get("error", "AssertionError / Unknown Error")
                if beta.get("traceback"):
                    failure_reason = f"{failure_reason}\n{beta.get('traceback')}"
                worksheet.append([
                    "No",
                    "N/A",
                    "N/A",
                    "ERR_FAILED",
                    "N/A",
                    "N/A",
                    failure_reason,
                    "",
                    "",
                    "",
                    "",
                    "",
                ])

        workbook.save(output_xlsx)

        if not output_xlsx.exists() or output_xlsx.stat().st_size == 0:
            logger.error("XLSX write verification failed: %s", output_xlsx)
            return 1

        logger.info("SUCCESS: Benchmarking complete. Report saved to: %s", output_xlsx)

        # Run automated regression gate check if possible
        if results:
            candidate_res = results[0]
            candidate_run_id = candidate_res.get("run_id")
            derived_bench = candidate_res.get("derived_benchmark")

            if candidate_run_id is not None and derived_bench:
                with SessionLocal() as session:
                    # Look for baseline labeled run
                    baseline_run = session.execute(
                        select(CalibrationRun)
                        .where(
                            CalibrationRun.benchmark_name == derived_bench,
                            CalibrationRun.success == True,
                            CalibrationRun.run_label == 'baseline'
                        )
                        .order_by(CalibrationRun.created_at.desc(), CalibrationRun.id.desc())
                    ).scalars().first()

                    # Fallback to any previous successful run for this benchmark
                    if not baseline_run:
                        baseline_run = session.execute(
                            select(CalibrationRun)
                            .where(
                                CalibrationRun.benchmark_name == derived_bench,
                                CalibrationRun.success == True,
                                CalibrationRun.id != candidate_run_id
                            )
                            .order_by(CalibrationRun.created_at.desc(), CalibrationRun.id.desc())
                        ).scalars().first()

                    if baseline_run:
                        logger.info("Comparing Candidate Run %s against Baseline Run %s", candidate_run_id, baseline_run.id)
                        passes, diff = regression_gate(
                            derived_bench,
                            baseline_run_id=baseline_run.id,
                            candidate_run_id=candidate_run_id,
                            allow_regressions=allow_regressions
                        )
                        logger.info("Scorecard Diff: %s", json.dumps(diff, indent=2))
                        if not passes:
                            logger.error("REGRESSION GATE FAILED: %s regressions found (allowed: %s)", len(diff["regressed"]), allow_regressions)
                            return 2
                        else:
                            logger.info("REGRESSION GATE PASSED: %s regressions found (allowed: %s)", len(diff["regressed"]), allow_regressions)
                    else:
                        logger.info("No previous baseline run found to compare for benchmark: %s", derived_bench)

        return 0
    finally:
        logger.removeHandler(file_handler)
        file_handler.close()


def main() -> int:
    parser = argparse.ArgumentParser(description="TM_BETA_TEST_LOCAL SQLAlchemy-backed PDF/DOCX runner.")
    parser.add_argument(
        "pdf_paths",
        type=Path,
        nargs="+",
        help="One or more paths to trademark search report PDF or DOCX files.",
    )
    parser.add_argument(
        "--output-xlsx",
        type=Path,
        help="Optional path to save the XLSX report. Defaults to TAEB_Calib/{first_pdf}_comparison.xlsx",
    )
    parser.add_argument(
        "--gt-json",
        type=str,
        help="Ground Truth JSON string to seed for the target PDF(s).",
    )
    parser.add_argument(
        "--gt-file",
        type=Path,
        help="Path to Ground Truth JSON file to seed for the target PDF(s).",
    )
    parser.add_argument(
        "--benchmark-name",
        type=str,
        help="Optional benchmark name to use as a stable identity.",
    )
    parser.add_argument(
        "--run-label",
        type=str,
        help="Optional label to categorize this run (e.g. baseline, candidate).",
    )
    parser.add_argument(
        "--git-commit",
        type=str,
        help="Optional git commit hash associated with this run.",
    )
    parser.add_argument(
        "--allow-regressions",
        type=int,
        default=0,
        help="Max regressions allowed before failing the regression gate (default: 0).",
    )
    args = parser.parse_args()

    if args.output_xlsx:
        output_xlsx = args.output_xlsx
    else:
        output_xlsx = OUTPUT_DIR / f"{args.pdf_paths[0].stem}_comparison.xlsx"

    return asyncio.run(
        main_async(
            args.pdf_paths,
            output_xlsx,
            gt_json=args.gt_json,
            gt_file=args.gt_file,
            benchmark_name=args.benchmark_name,
            run_label=args.run_label,
            git_commit=args.git_commit,
            allow_regressions=args.allow_regressions,
        )
    )



if __name__ == "__main__":
    sys.exit(main())
